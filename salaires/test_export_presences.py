from datetime import date, time
from decimal import Decimal
from io import BytesIO

from django.contrib.auth import get_user_model
from django.test import TestCase, override_settings
from django.urls import reverse
from openpyxl import load_workbook

from eleves.models import Ecole
from .models import Enseignant, PresenceEnseignant, TypeEnseignant
from .tests import TEST_MIDDLEWARE


@override_settings(MIDDLEWARE=TEST_MIDDLEWARE)
class ExportPresencesExcelTests(TestCase):
    @classmethod
    def setUpTestData(cls):
        cls.ecole = Ecole.objects.create(nom="École A", adresse="Conakry", telephone="620000001", directeur="Test")
        cls.autre_ecole = Ecole.objects.create(nom="École B", adresse="Conakry", telephone="620000002", directeur="Test")
        cls.user = get_user_model().objects.create_user(username="export-presences")
        cls.user.profil.ecole = cls.ecole
        cls.user.profil.role = "ADMIN"
        cls.user.profil.is_validated = True
        cls.user.profil.save()
        cls.admin = get_user_model().objects.create_superuser(username="export-superadmin", email="test@example.com")
        cls.admin.profil.ecole = None
        cls.admin.profil.save()
        cls.enseignant = Enseignant.objects.create(
            ecole=cls.ecole, nom="Diallo", prenoms="Test", type_enseignant=TypeEnseignant.CHAUFFEUR,
            salaire_fixe=Decimal("1000000"), date_embauche=date(2026, 1, 1),
        )
        cls.autre_enseignant = Enseignant.objects.create(
            ecole=cls.autre_ecole, nom="Camara", prenoms="Test", type_enseignant=TypeEnseignant.PRIMAIRE,
            salaire_fixe=Decimal("1000000"), date_embauche=date(2026, 1, 1),
        )
        for jour, statut, heures in ((1, "PRESENT", "7.5"), (2, "ABSENT", "0"), (7, "RETARD", "6"), (8, "PRESENT", "8")):
            PresenceEnseignant.objects.create(
                enseignant=cls.enseignant, date=date(2026, 9, jour), statut=statut,
                heures_travaillees=Decimal(heures), observations="Pointage de test", pointe_par=cls.user,
            )
        PresenceEnseignant.objects.create(
            enseignant=cls.autre_enseignant, date=date(2026, 9, 3),
            statut="PRESENT", heure_arrivee=time(8), heure_depart=time(16), pointe_par=cls.user,
        )

    def exporter(self, user=None, **filtres):
        self.client.force_login(user or self.user)
        response = self.client.get(reverse("salaires:export_presences_excel"), {
            "date_debut": "2026-09-01", "date_fin": "2026-09-07", **filtres,
        })
        self.assertEqual(response.status_code, 200)
        self.assertEqual(response["Content-Type"], "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")
        self.assertIn("presences_20260901_20260907.xlsx", response["Content-Disposition"])
        workbook = load_workbook(BytesIO(response.content))
        self.addCleanup(workbook.close)
        return workbook

    def test_periode_demandee_et_totaux_sans_autre_ecole(self):
        workbook = self.exporter()
        detail = workbook["Détail Présences"]
        self.assertIn(self.ecole.nom, detail["A1"].value)
        self.assertEqual([detail.cell(row, 1).value for row in range(5, detail.max_row + 1)],
                         ["01/09/2026", "02/09/2026", "07/09/2026"])
        self.assertEqual(detail["C5"].value, "Chauffeur")
        self.assertEqual(detail["G6"].value, 0)
        recap = workbook["Récapitulatif"]
        self.assertEqual([recap.cell(4, col).value for col in range(3, 10)], [1, 1, 1, 0, 0, 0, 13.5])
        self.assertEqual(recap["A5"].value, "TOTAL")
        self.assertEqual(recap["I5"].value, 13.5)

    def test_superadmin_sans_ecole_peut_exporter(self):
        workbook = self.exporter(self.admin)
        self.assertIn("Toutes les écoles", workbook["Détail Présences"]["A1"].value)
        self.assertEqual(workbook["Détail Présences"].max_row, 8)
        self.assertEqual(workbook["Récapitulatif"]["I6"].value, 21.5)

    def test_superadmin_sans_profil_peut_exporter(self):
        self.admin.profil.delete()
        workbook = self.exporter(self.admin)
        self.assertEqual(workbook["Détail Présences"].max_row, 8)

    def test_compte_sans_ecole_ne_divulgue_aucune_presence(self):
        self.user.profil.ecole = None
        self.user.profil.save()
        workbook = self.exporter()
        self.assertEqual(workbook["Détail Présences"].max_row, 4)
        self.assertEqual(workbook["Récapitulatif"]["I4"].value, 0)

    def test_filtre_enseignant_autre_ecole_reste_vide(self):
        workbook = self.exporter(enseignant=self.autre_enseignant.pk)
        self.assertEqual(workbook["Détail Présences"].max_row, 4)
        self.assertEqual(workbook["Récapitulatif"]["I4"].value, 0)

    def test_export_vide_reste_un_classeur_valide(self):
        PresenceEnseignant.objects.all().delete()
        workbook = self.exporter()
        self.assertEqual(workbook.sheetnames, ["Détail Présences", "Récapitulatif"])
        self.assertEqual(workbook["Récapitulatif"]["I4"].value, 0)

    def test_superadmin_avec_ecole_conserve_le_perimetre(self):
        self.admin.profil.ecole = self.ecole
        self.admin.profil.save()
        workbook = self.exporter(self.admin)
        self.assertIn(self.ecole.nom, workbook["Détail Présences"]["A1"].value)
        self.assertEqual(workbook["Détail Présences"].max_row, 7)
        self.assertEqual(workbook["Récapitulatif"]["I5"].value, 13.5)
