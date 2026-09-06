from datetime import date
from io import BytesIO

from django.contrib.auth import get_user_model
from django.test import TestCase, override_settings
from django.urls import reverse
from openpyxl import load_workbook

from eleves.models import Classe, Ecole, Eleve
from paiements.models import EcheancierPaiement, ModePaiement, Paiement, PaiementRemise, RemiseReduction, TypePaiement
from .support import TEST_MIDDLEWARE


@override_settings(MIDDLEWARE=TEST_MIDDLEWARE)
class RecapClasseVersementsTests(TestCase):
    def setUp(self):
        self.ecole = Ecole.objects.create(nom='École récapitulatif', adresse='Conakry', telephone='620000001', directeur='Direction')
        self.classe_a = Classe.objects.create(ecole=self.ecole, nom='CP A', niveau='PRIMAIRE_1', annee_scolaire='2025-2026')
        self.classe_b = Classe.objects.create(ecole=self.ecole, nom='CP B', niveau='PRIMAIRE_1', annee_scolaire='2025-2026')
        self.eleves = []
        for index, classe in enumerate((self.classe_a, self.classe_a, self.classe_b)):
            eleve = Eleve.objects.create(matricule=f'RECAP-{index}', nom=f'Élève {index}', prenom='Test', sexe='F', classe=classe)
            EcheancierPaiement.objects.create(
                eleve=eleve, annee_scolaire='2025-2026', frais_inscription_du=30000,
                tranche_1_due=100000, tranche_2_due=100000, tranche_3_due=100000,
                date_echeance_inscription=date(2025, 9, 1), date_echeance_tranche_1=date(2026, 1, 15),
                date_echeance_tranche_2=date(2026, 3, 15), date_echeance_tranche_3=date(2026, 5, 15),
            )
            self.eleves.append(eleve)
        self.type_paiement = TypePaiement.objects.create(nom='Scolarité annuelle')
        self.mode = ModePaiement.objects.create(nom='Espèces')
        self.premier = self.payer(self.eleves[0])
        self.deuxieme = self.payer(self.eleves[0])
        self.payer(self.eleves[1])
        self.user = get_user_model().objects.create_superuser('recap-versements', 'test@example.com', 'test')
        self.client.force_login(self.user)
        self.url = reverse('paiements:liste_paiements')

    def payer(self, eleve, **extra):
        data = dict(eleve=eleve, type_paiement=self.type_paiement, mode_paiement=self.mode,
                    montant=50000, statut='VALIDE', date_paiement=date(2025, 10, 1))
        data.update(extra)
        return Paiement.objects.create(**data)

    def remise(self, paiement, montant):
        remise = RemiseReduction.objects.create(
            nom=f'Remise {PaiementRemise.objects.count()}', type_remise='MONTANT_FIXE', valeur=montant,
            motif='AUTRE', date_debut=date(2025, 9, 1), date_fin=date(2026, 7, 1),
        )
        return PaiementRemise.objects.create(paiement=paiement, remise=remise, montant_remise=montant, portee_tranches='1')

    def verifier_recap(self, reduction=0):
        response = self.client.get(self.url)
        self.assertEqual(response.status_code, 200)
        totaux = response.context['totaux_du']
        self.assertEqual(totaux['eleves_count'], 3)
        self.assertEqual(totaux['du_sco_net'], 900000 - reduction)
        self.assertEqual(totaux['frais_inscription_total'], 90000)
        self.assertEqual(totaux['du_global_net'], 990000 - reduction)
        classes = {row['classe_id']: row for row in response.context['totaux_du_detail_classes']}
        self.assertEqual(classes[self.classe_a.pk]['eleves_count'], 2)
        self.assertEqual(classes[self.classe_a.pk]['du_sco_net'], 600000 - reduction)
        self.assertEqual(classes[self.classe_a.pk]['frais_inscription_total'], 60000)
        self.assertEqual(classes[self.classe_a.pk]['du_global_net'], 660000 - reduction)
        self.assertEqual(classes[self.classe_b.pk]['du_global_net'], 330000)
        self.assertEqual(sum(row['du_global_net'] for row in classes.values()), totaux['du_global_net'])
        return classes

    def verifier_excel(self, classes):
        response = self.client.get(reverse('paiements:export_recap_par_classe_excel'))
        self.assertEqual(response.status_code, 200)
        workbook = load_workbook(BytesIO(response.content), read_only=True)
        lignes = {row[1]: row for row in workbook.active.iter_rows(min_row=2, values_only=True)}
        for row in classes.values():
            export = lignes[row['classe_nom']]
            self.assertEqual(export[2], row['eleves_count'])
            self.assertEqual(export[3], row['du_sco_net'])
            self.assertEqual(export[4], row['frais_inscription_total'])
            self.assertEqual(export[7], row['du_global_net'])
        workbook.close()

    def test_plusieurs_versements_comptent_un_seul_du_par_eleve(self):
        self.verifier_excel(self.verifier_recap())

    def test_plusieurs_remises_identiques_ne_multiplient_pas_les_frais(self):
        self.remise(self.premier, 20000)
        self.remise(self.premier, 20000)
        self.remise(self.deuxieme, 10000)
        self.verifier_excel(self.verifier_recap(reduction=50000))

    def test_modification_suppression_et_statut_ne_changent_pas_le_du(self):
        self.premier.montant = 60000
        self.premier.save()
        self.deuxieme.delete()
        self.payer(self.eleves[0], statut='EN_ATTENTE')
        self.payer(self.eleves[0], statut='ANNULE')
        self.verifier_recap()

    def test_remises_hors_annee_ecole_ou_non_validees_sont_exclues(self):
        self.remise(self.premier, 10000)
        autre_ecole = Ecole.objects.create(nom='École externe', adresse='Conakry', telephone='620000002', directeur='Direction')
        for extra in ({'annee_scolaire': '2024-2025'}, {'ecole_encaissement': autre_ecole},
                      {'statut': 'EN_ATTENTE'}, {'statut': 'ANNULE'}):
            self.remise(self.payer(self.eleves[0], **extra), 50000)
        self.verifier_recap(reduction=10000)

    def test_recap_filtre_classe_identique_en_ajax(self):
        for headers in ({}, {'HTTP_X_REQUESTED_WITH': 'XMLHttpRequest'}):
            with self.subTest(headers=headers):
                response = self.client.get(self.url, {'classe': self.classe_a.pk}, **headers)
                self.assertEqual(response.status_code, 200)
                self.assertEqual(response.context['totaux_du']['eleves_count'], 2)
                self.assertEqual(response.context['totaux_du']['du_global_net'], 660000)
                rows = response.context['totaux_du_detail_classes']
                self.assertEqual(len(rows), 1)
                self.assertEqual(rows[0]['du_global_net'], 660000)
