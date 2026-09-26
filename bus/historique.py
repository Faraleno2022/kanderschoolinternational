"""Historique des abonnements d'un élève (bus et cantine).

- reprise des informations du dernier abonnement lors d'un nouvel abonnement ;
- liste des abonnements d'un élève, exportable en Excel et en PDF ;
- carnet d'abonnement PDF (mois, montant payé, dates, jour d'expiration).
"""
import io
import os
from datetime import timedelta

from dateutil.relativedelta import relativedelta
from django.contrib.auth.decorators import login_required
from django.http import Http404, HttpResponse, JsonResponse
from django.shortcuts import get_object_or_404, render
from django.urls import reverse
from django.utils import timezone
from django.utils.html import escape

from eleves.models import Eleve
from utilisateurs.utils import filter_by_user_school, user_is_superadmin

from .models import AbonnementBus, AbonnementCantine, TypePeriodiciteAbonnement

MOIS = [
    'Janvier', 'Février', 'Mars', 'Avril', 'Mai', 'Juin', 'Juillet',
    'Août', 'Septembre', 'Octobre', 'Novembre', 'Décembre',
]
JOURS = ['Lundi', 'Mardi', 'Mercredi', 'Jeudi', 'Vendredi', 'Samedi', 'Dimanche']

SERVICES = {
    'bus': {
        'model': AbonnementBus,
        'related': 'abonnements_bus',
        'service_periodicite': TypePeriodiciteAbonnement.Service.BUS,
        'titre': 'Bus scolaire',
        'champs_repris': [
            'montant', 'periodicite', 'alerte_avant_jours',
            'zone', 'itineraire', 'point_arret', 'contact_parent',
        ],
        'url_nouveau': 'bus:nouveau',
        'url_liste': 'bus:liste',
        'url_modifier': 'bus:modifier',
        'url_recu': 'bus:recu_pdf',
    },
    'cantine': {
        'model': AbonnementCantine,
        'related': 'abonnements_cantine',
        'service_periodicite': TypePeriodiciteAbonnement.Service.CANTINE,
        'titre': 'Cantine scolaire',
        'champs_repris': [
            'montant', 'periodicite', 'type_repas', 'alerte_avant_jours',
            'regime_alimentaire', 'allergies', 'contact_parent',
        ],
        'url_nouveau': 'bus:creer_abonnement_cantine',
        'url_liste': 'bus:liste_abonnements_cantine',
        'url_modifier': 'bus:modifier_abonnement_cantine',
        'url_recu': 'bus:recu_cantine_pdf',
    },
}

DUREES_PAR_DEFAUT = {
    'JOURNALIER': (0, 1), 'HEBDOMADAIRE': (0, 7), 'MENSUEL': (1, 0),
    'TRIMESTRIEL': (3, 0), 'SEMESTRIEL': (6, 0), 'ANNUEL': (12, 0),
}


# ---------------------------------------------------------------------------
# Utilitaires
# ---------------------------------------------------------------------------

def _config(service):
    config = SERVICES.get(service)
    if config is None:
        raise Http404("Service d'abonnement inconnu.")
    return config


def _eleve_autorise(user, eleve_id):
    qs = Eleve.objects.select_related('classe', 'classe__ecole', 'responsable_principal')
    if not user_is_superadmin(user):
        qs = filter_by_user_school(qs, user, 'classe__ecole')
    return get_object_or_404(qs, pk=eleve_id)


def _abonnements(eleve, service):
    return getattr(eleve, _config(service)['related']).order_by('date_debut', 'id')


def libelle_mois(abonnement):
    """Mois couverts : « Septembre 2026 » ou « Septembre – Novembre 2026 »."""
    debut, fin = abonnement.date_debut, abonnement.date_expiration
    if not debut:
        return '-'
    libelle = f"{MOIS[debut.month - 1]} {debut.year}"
    if fin and (fin - debut).days > 31 and (fin.year, fin.month) != (debut.year, debut.month):
        if fin.year == debut.year:
            libelle = f"{MOIS[debut.month - 1]} – {MOIS[fin.month - 1]} {fin.year}"
        else:
            libelle = f"{MOIS[debut.month - 1]} {debut.year} – {MOIS[fin.month - 1]} {fin.year}"
    return libelle


def jour_semaine(valeur):
    return JOURS[valeur.weekday()] if valeur else '-'


def _statut(abonnement, today):
    if abonnement.date_expiration and abonnement.date_expiration < today:
        return 'Expiré'
    return abonnement.get_statut_display()


def _detail_service(abonnement, service):
    if service == 'cantine':
        return abonnement.get_type_repas_display()
    morceaux = [abonnement.zone, abonnement.point_arret]
    return ' / '.join(m for m in morceaux if m) or '-'


def lignes_historique(eleve, service):
    today = timezone.localdate()
    lignes = []
    for abonnement in _abonnements(eleve, service):
        jours_restants = (
            (abonnement.date_expiration - today).days if abonnement.date_expiration else None
        )
        lignes.append({
            'abonnement': abonnement,
            'mois': libelle_mois(abonnement),
            'periodicite': abonnement.get_periodicite_display(),
            'detail': _detail_service(abonnement, service),
            'montant': abonnement.montant or 0,
            'reference': abonnement.reference_paiement or '',
            'date_debut': abonnement.date_debut,
            'date_expiration': abonnement.date_expiration,
            'jour_expiration': jour_semaine(abonnement.date_expiration),
            'jours_restants': jours_restants,
            'statut': _statut(abonnement, today),
        })
    return lignes


def _duree(service, periodicite):
    option = TypePeriodiciteAbonnement.objects.filter(
        service=_config(service)['service_periodicite'], code=periodicite,
    ).first()
    if option and (option.duree_mois or option.duree_jours):
        return option.duree_mois, option.duree_jours
    return DUREES_PAR_DEFAUT.get(periodicite, (0, 0))


def valeurs_reprises(eleve, service):
    """Valeurs du dernier abonnement de l'élève, pour pré-remplir le suivant.

    La date de début proposée est le lendemain de la dernière expiration,
    ou aujourd'hui si l'abonnement précédent est déjà expiré.
    """
    dernier = getattr(eleve, _config(service)['related']).order_by('-date_expiration', '-id').first()
    if dernier is None:
        return None
    valeurs = {champ: getattr(dernier, champ) for champ in _config(service)['champs_repris']}
    today = timezone.localdate()
    debut = today
    if dernier.date_expiration:
        debut = max(dernier.date_expiration + timedelta(days=1), today)
    mois, jours = _duree(service, dernier.periodicite)
    if mois or jours:
        expiration = debut + relativedelta(months=mois, days=jours)
    elif dernier.date_debut and dernier.date_expiration:
        expiration = debut + (dernier.date_expiration - dernier.date_debut)
    else:
        expiration = None
    valeurs['date_debut'] = debut
    valeurs['date_expiration'] = expiration
    return {'valeurs': valeurs, 'dernier': dernier}


# ---------------------------------------------------------------------------
# Vues
# ---------------------------------------------------------------------------

@login_required
def dernier_abonnement_json(request, service, eleve_id):
    """Informations reprises du dernier abonnement (pré-remplissage du formulaire)."""
    eleve = _eleve_autorise(request.user, eleve_id)
    reprise = valeurs_reprises(eleve, service)
    data = {
        'success': True,
        'existe': reprise is not None,
        'nombre': _abonnements(eleve, service).count(),
        'historique_url': reverse('bus:historique_eleve', args=[eleve.pk, service]),
    }
    if reprise:
        valeurs = {
            cle: (valeur.isoformat() if hasattr(valeur, 'isoformat') else ('' if valeur is None else str(valeur)))
            for cle, valeur in reprise['valeurs'].items()
        }
        dernier = reprise['dernier']
        data.update({
            'valeurs': valeurs,
            'dernier': {
                'date_debut': dernier.date_debut.strftime('%d/%m/%Y') if dernier.date_debut else '',
                'date_expiration': dernier.date_expiration.strftime('%d/%m/%Y') if dernier.date_expiration else '',
            },
        })
    return JsonResponse(data)


@login_required
def historique_eleve(request, service, eleve_id):
    config = _config(service)
    eleve = _eleve_autorise(request.user, eleve_id)
    lignes = lignes_historique(eleve, service)
    return render(request, 'bus/historique_eleve.html', {
        'titre_page': f"Abonnements {config['titre'].lower()} — {eleve.prenom} {eleve.nom}",
        'eleve': eleve,
        'service': service,
        'config': config,
        'lignes': lignes,
        'total': sum(ligne['montant'] for ligne in lignes),
    })


@login_required
def export_historique_excel(request, service, eleve_id):
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter

    config = _config(service)
    eleve = _eleve_autorise(request.user, eleve_id)
    lignes = lignes_historique(eleve, service)

    wb = Workbook()
    ws = wb.active
    ws.title = 'Abonnements'
    ws.append([f"Abonnements {config['titre'].lower()} — {eleve.prenom} {eleve.nom}"])
    ws['A1'].font = Font(bold=True, size=14)
    ws.append([f"Matricule : {eleve.matricule or '-'}   Classe : {getattr(eleve.classe, 'nom', '-')}"
               f"   École : {getattr(eleve.classe.ecole, 'nom', '-')}"])
    ws.append([])

    entetes = [
        'Mois', 'Type', 'Repas' if service == 'cantine' else 'Zone / Arrêt', 'Montant payé (GNF)',
        'Référence', 'Date début', "Date d'expiration", "Jour d'expiration", 'Jours restants', 'Statut',
    ]
    ws.append(entetes)
    for cellule in ws[4]:
        cellule.font = Font(bold=True, color='FFFFFF')
        cellule.fill = PatternFill('solid', fgColor='1746A2')
        cellule.alignment = Alignment(horizontal='center', vertical='center', wrap_text=True)

    for ligne in lignes:
        ws.append([
            ligne['mois'], ligne['periodicite'], ligne['detail'], float(ligne['montant']),
            ligne['reference'], ligne['date_debut'], ligne['date_expiration'],
            ligne['jour_expiration'], ligne['jours_restants'], ligne['statut'],
        ])
        row = ws.max_row
        ws.cell(row, 4).number_format = '#,##0'
        ws.cell(row, 6).number_format = 'DD/MM/YYYY'
        ws.cell(row, 7).number_format = 'DD/MM/YYYY'

    ws.append([])
    ws.append(['TOTAL PAYÉ', '', '', float(sum(ligne['montant'] for ligne in lignes))])
    ws.cell(ws.max_row, 1).font = Font(bold=True)
    ws.cell(ws.max_row, 4).font = Font(bold=True)
    ws.cell(ws.max_row, 4).number_format = '#,##0'

    for index, largeur in enumerate([24, 16, 22, 18, 20, 13, 16, 16, 14, 12], start=1):
        ws.column_dimensions[get_column_letter(index)].width = largeur

    buffer = io.BytesIO()
    wb.save(buffer)
    response = HttpResponse(
        buffer.getvalue(),
        content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet',
    )
    response['Content-Disposition'] = (
        f'attachment; filename="abonnements_{service}_{eleve.matricule or eleve.pk}.xlsx"'
    )
    return response


# ---------------------------------------------------------------------------
# PDF
# ---------------------------------------------------------------------------

def _polices():
    from reportlab.pdfbase import pdfmetrics
    from reportlab.pdfbase.ttfonts import TTFont
    try:
        pdfmetrics.registerFont(TTFont('Arial', 'C:/Windows/Fonts/arial.ttf'))
        pdfmetrics.registerFont(TTFont('Arial-Bold', 'C:/Windows/Fonts/arialbd.ttf'))
        return 'Arial', 'Arial-Bold'
    except Exception:
        return 'Helvetica', 'Helvetica-Bold'


def _fichier(champ):
    try:
        if champ and hasattr(champ, 'path') and os.path.exists(champ.path):
            return champ.path
    except Exception:
        pass
    return None


def _gnf(valeur):
    return f"{int(valeur or 0):,}".replace(',', ' ') + ' GNF'


def _entete(eleve, titre, largeur, police, police_gras):
    """Bloc d'en-tête : logo, nom de l'école, titre du document."""
    from reportlab.lib import colors
    from reportlab.lib.styles import ParagraphStyle
    from reportlab.platypus import Image, Paragraph, Table, TableStyle

    ecole = eleve.classe.ecole
    style_ecole = ParagraphStyle('ecole', fontName=police_gras, fontSize=15, leading=18,
                                 textColor=colors.HexColor('#1746a2'))
    style_titre = ParagraphStyle('titre', fontName=police_gras, fontSize=11, leading=14,
                                 textColor=colors.HexColor('#0f766e'))
    style_info = ParagraphStyle('info', fontName=police, fontSize=8.5, leading=11,
                                textColor=colors.HexColor('#64748b'))
    infos = ' — '.join(x for x in [getattr(ecole, 'adresse', ''), getattr(ecole, 'telephone', '')] if x)
    texte = [
        Paragraph(escape((ecole.nom or '').upper()), style_ecole),
        Paragraph(titre, style_titre),
        Paragraph(f"Année scolaire {eleve.classe.annee_scolaire or '-'}"
                  + (f" — {escape(infos)}" if infos else ''), style_info),
    ]
    logo = _fichier(getattr(ecole, 'logo', None))
    cellule_logo = Image(logo, width=52, height=52, kind='proportional') if logo else ''
    table = Table([[cellule_logo, texte]], colWidths=[62, largeur - 62])
    table.setStyle(TableStyle([
        ('VALIGN', (0, 0), (-1, -1), 'MIDDLE'),
        ('LINEBELOW', (0, 0), (-1, 0), 1.2, colors.HexColor('#1746a2')),
        ('BOTTOMPADDING', (0, 0), (-1, -1), 8),
    ]))
    return table


def _style_tableau(entete_couleur='#1746a2'):
    from reportlab.lib import colors
    return [
        ('BACKGROUND', (0, 0), (-1, 0), colors.HexColor(entete_couleur)),
        ('TEXTCOLOR', (0, 0), (-1, 0), colors.white),
        ('ALIGN', (0, 0), (-1, -1), 'CENTER'),
        ('VALIGN', (0, 0), (-1, -1), 'MIDDLE'),
        ('GRID', (0, 0), (-1, -1), 0.5, colors.HexColor('#cbd5e1')),
        ('ROWBACKGROUNDS', (0, 1), (-1, -1), [colors.white, colors.HexColor('#f5f8fc')]),
        ('TOPPADDING', (0, 0), (-1, -1), 4),
        ('BOTTOMPADDING', (0, 0), (-1, -1), 4),
    ]


def _date(valeur):
    return valeur.strftime('%d/%m/%Y') if valeur else '-'


@login_required
def export_historique_pdf(request, service, eleve_id):
    """Liste des abonnements de l'élève (A4 paysage)."""
    from reportlab.lib import colors
    from reportlab.lib.pagesizes import A4, landscape
    from reportlab.lib.styles import ParagraphStyle
    from reportlab.lib.units import mm
    from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle

    config = _config(service)
    eleve = _eleve_autorise(request.user, eleve_id)
    lignes = lignes_historique(eleve, service)
    police, police_gras = _polices()

    response = HttpResponse(content_type='application/pdf')
    response['Content-Disposition'] = (
        f'attachment; filename="abonnements_{service}_{eleve.matricule or eleve.pk}.pdf"'
    )
    doc = SimpleDocTemplate(response, pagesize=landscape(A4), leftMargin=12 * mm, rightMargin=12 * mm,
                            topMargin=10 * mm, bottomMargin=10 * mm,
                            title=f"Abonnements {config['titre']} - {eleve.prenom} {eleve.nom}")
    largeur = doc.width
    style = ParagraphStyle('p', fontName=police, fontSize=9.5, leading=13)

    elements = [
        _entete(eleve, f"LISTE DES ABONNEMENTS — {config['titre'].upper()}", largeur, police, police_gras),
        Spacer(1, 6),
        Paragraph(
            f"<b>Élève :</b> {escape(eleve.prenom)} {escape(eleve.nom)} &nbsp;&nbsp; "
            f"<b>Matricule :</b> {escape(eleve.matricule or '-')} &nbsp;&nbsp; "
            f"<b>Classe :</b> {escape(getattr(eleve.classe, 'nom', '-'))}", style,
        ),
        Spacer(1, 8),
    ]
    cellule = ParagraphStyle('cellule', fontName=police, fontSize=8.5, leading=10, alignment=1)
    donnees = [[
        'Mois', 'Type', 'Repas' if service == 'cantine' else 'Zone / Arrêt', 'Montant\npayé',
        'Référence', 'Date\ndébut', "Date\nd'expiration", "Jour\nd'expiration", 'Jours\nrestants', 'Statut',
    ]]
    for ligne in lignes:
        donnees.append([
            Paragraph(escape(ligne['mois']), cellule), ligne['periodicite'],
            Paragraph(escape(ligne['detail']), cellule), _gnf(ligne['montant']),
            ligne['reference'] or '-', _date(ligne['date_debut']), _date(ligne['date_expiration']),
            ligne['jour_expiration'],
            '-' if ligne['jours_restants'] is None else str(ligne['jours_restants']),
            ligne['statut'],
        ])
    if not lignes:
        donnees.append(['Aucun abonnement'] + [''] * 9)
    donnees.append(['TOTAL PAYÉ', '', '', _gnf(sum(ligne['montant'] for ligne in lignes))] + [''] * 6)

    ratios = [0.15, 0.09, 0.13, 0.10, 0.11, 0.08, 0.09, 0.09, 0.07, 0.09]
    table = Table(donnees, colWidths=[largeur * r for r in ratios], repeatRows=1)
    table.setStyle(TableStyle(_style_tableau() + [
        ('FONTNAME', (0, 0), (-1, -1), police),
        ('FONTNAME', (0, 0), (-1, 0), police_gras),
        ('FONTNAME', (0, -1), (-1, -1), police_gras),
        ('FONTSIZE', (0, 0), (-1, -1), 8.5),
        ('SPAN', (0, -1), (2, -1)),
        ('BACKGROUND', (0, -1), (-1, -1), colors.HexColor('#e2e8f0')),
    ]))
    elements.append(table)
    elements.append(Spacer(1, 8))
    elements.append(Paragraph(
        f"Édité le {timezone.localtime().strftime('%d/%m/%Y à %H:%M')}",
        ParagraphStyle('pied', fontName=police, fontSize=8, textColor=colors.HexColor('#64748b')),
    ))
    doc.build(elements)
    return response


@login_required
def carnet_abonnement_pdf(request, service, eleve_id):
    """Carnet d'abonnement de l'élève : un abonnement par ligne, avec cases de visa."""
    from reportlab.lib import colors
    from reportlab.lib.pagesizes import A4
    from reportlab.lib.styles import ParagraphStyle
    from reportlab.lib.units import mm
    from reportlab.platypus import Image, Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle

    config = _config(service)
    eleve = _eleve_autorise(request.user, eleve_id)
    lignes = lignes_historique(eleve, service)
    police, police_gras = _polices()

    response = HttpResponse(content_type='application/pdf')
    response['Content-Disposition'] = (
        f'attachment; filename="carnet_{service}_{eleve.matricule or eleve.pk}.pdf"'
    )
    doc = SimpleDocTemplate(response, pagesize=A4, leftMargin=14 * mm, rightMargin=14 * mm,
                            topMargin=12 * mm, bottomMargin=12 * mm,
                            title=f"Carnet d'abonnement {config['titre']} - {eleve.prenom} {eleve.nom}")
    largeur = doc.width
    label = ParagraphStyle('label', fontName=police_gras, fontSize=8, leading=10,
                           textColor=colors.HexColor('#64748b'))
    valeur = ParagraphStyle('valeur', fontName=police_gras, fontSize=10.5, leading=13,
                            textColor=colors.HexColor('#0f172a'))

    # Identité de l'élève et informations du service (dernier abonnement)
    dernier = lignes[-1]['abonnement'] if lignes else None
    responsable = eleve.responsable_principal
    contact = (getattr(dernier, 'contact_parent', '') if dernier else '') or getattr(responsable, 'telephone', '') or '-'
    champs = [
        ('ÉLÈVE', f"{eleve.prenom} {eleve.nom}".upper()),
        ('MATRICULE', eleve.matricule or '-'),
        ('CLASSE', getattr(eleve.classe, 'nom', '-')),
        ('PARENT / CONTACT', f"{getattr(responsable, 'prenom', '')} {getattr(responsable, 'nom', '')}".strip()
         + (f" — {contact}" if contact != '-' else '') if responsable else contact),
    ]
    if service == 'cantine':
        champs += [
            ('TYPE DE REPAS', dernier.get_type_repas_display() if dernier else '-'),
            ('RÉGIME / ALLERGIES', ' — '.join(x for x in [
                getattr(dernier, 'regime_alimentaire', ''), getattr(dernier, 'allergies', '')] if x) or '-'),
        ]
    else:
        champs += [
            ('ZONE / ITINÉRAIRE', ' — '.join(x for x in [
                getattr(dernier, 'zone', ''), getattr(dernier, 'itineraire', '')] if x) or '-'),
            ("POINT D'ARRÊT", getattr(dernier, 'point_arret', '') or '-' if dernier else '-'),
        ]
    grille = []
    for i in range(0, len(champs), 2):
        rangee = []
        for lib, val in champs[i:i + 2]:
            rangee.append([Paragraph(escape(lib), label), Paragraph(escape(str(val)), valeur)])
        grille.append(rangee)
    photo = _fichier(getattr(eleve, 'photo', None))
    largeur_photo = 30 * mm
    identite = Table(grille, colWidths=[(largeur - largeur_photo - 6) / 2] * 2)
    identite.setStyle(TableStyle([('VALIGN', (0, 0), (-1, -1), 'TOP'),
                                  ('BOTTOMPADDING', (0, 0), (-1, -1), 5)]))
    cellule_photo = (Image(photo, width=largeur_photo - 4, height=36 * mm, kind='proportional')
                     if photo else Paragraph('PHOTO', ParagraphStyle(
                         'ph', fontName=police_gras, fontSize=9, alignment=1,
                         textColor=colors.HexColor('#94a3b8'))))
    bloc_identite = Table([[identite, cellule_photo]], colWidths=[largeur - largeur_photo, largeur_photo],
                          rowHeights=[40 * mm])
    bloc_identite.setStyle(TableStyle([
        ('VALIGN', (0, 0), (-1, -1), 'MIDDLE'),
        ('ALIGN', (1, 0), (1, 0), 'CENTER'),
        ('BOX', (0, 0), (-1, -1), 0.8, colors.HexColor('#cbd5e1')),
        ('BOX', (1, 0), (1, 0), 0.8, colors.HexColor('#cbd5e1')),
        ('BACKGROUND', (0, 0), (0, 0), colors.HexColor('#f5f8fc')),
    ]))

    # Tableau du carnet (complété par des lignes vides à remplir à la main)
    mois_style = ParagraphStyle('mois', fontName=police_gras, fontSize=9, leading=10.5, alignment=1)
    donnees = [['N°', 'Mois', 'Montant\npayé', 'Date\ndébut', "Date\nd'expiration", "Jour\nd'expiration", 'Visa']]
    for numero, ligne in enumerate(lignes, start=1):
        donnees.append([
            str(numero), Paragraph(escape(ligne['mois']), mois_style), _gnf(ligne['montant']),
            _date(ligne['date_debut']),
            _date(ligne['date_expiration']), ligne['jour_expiration'], '',
        ])
    for numero in range(len(lignes) + 1, max(12, len(lignes)) + 1):
        donnees.append([str(numero), '', '', '', '', '', ''])
    donnees.append(['', 'TOTAL PAYÉ', _gnf(sum(ligne['montant'] for ligne in lignes)), '', '', '', ''])

    ratios = [0.06, 0.24, 0.16, 0.13, 0.15, 0.13, 0.13]
    carnet = Table(donnees, colWidths=[largeur * r for r in ratios], repeatRows=1,
                   rowHeights=[10 * mm] + [9 * mm] * (len(donnees) - 1))
    carnet.setStyle(TableStyle(_style_tableau('#0f766e') + [
        ('FONTNAME', (0, 0), (-1, -1), police),
        ('FONTNAME', (0, 0), (-1, 0), police_gras),
        ('FONTNAME', (0, -1), (-1, -1), police_gras),
        ('FONTSIZE', (0, 0), (-1, -1), 9),
        ('FONTNAME', (1, 1), (1, -1), police_gras),
        ('BACKGROUND', (0, -1), (-1, -1), colors.HexColor('#e2e8f0')),
    ]))

    signature = Table(
        [['Signature du parent', '', "Cachet et signature de l'administration"], ['', '', '']],
        colWidths=[largeur * 0.4, largeur * 0.2, largeur * 0.4], rowHeights=[6 * mm, 18 * mm],
    )
    signature.setStyle(TableStyle([
        ('FONTNAME', (0, 0), (-1, -1), police_gras),
        ('FONTSIZE', (0, 0), (-1, -1), 9),
        ('ALIGN', (0, 0), (-1, -1), 'CENTER'),
        ('LINEBELOW', (0, 1), (0, 1), 0.8, colors.HexColor('#475569')),
        ('LINEBELOW', (2, 1), (2, 1), 0.8, colors.HexColor('#475569')),
    ]))

    elements = [
        _entete(eleve, f"CARNET D'ABONNEMENT — {config['titre'].upper()}", largeur, police, police_gras),
        Spacer(1, 8),
        bloc_identite,
        Spacer(1, 10),
        carnet,
        Spacer(1, 12),
        signature,
        Spacer(1, 6),
        Paragraph(
            f"Édité le {timezone.localtime().strftime('%d/%m/%Y à %H:%M')} — "
            "Ce carnet doit être présenté à chaque paiement.",
            ParagraphStyle('pied', fontName=police, fontSize=8, textColor=colors.HexColor('#64748b')),
        ),
    ]
    doc.build(elements)
    return response
