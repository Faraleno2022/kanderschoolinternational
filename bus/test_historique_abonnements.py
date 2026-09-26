from datetime import date, timedelta
from io import BytesIO

from django.contrib.auth.models import User
from django.test import TestCase, override_settings
from django.urls import reverse
from django.utils import timezone
from openpyxl import load_workbook

from eleves.models import Classe, Ecole, Eleve
from paiements.tests.support import TEST_MIDDLEWARE

from .historique import jour_semaine, libelle_mois
from .models import AbonnementBus, AbonnementCantine


def _utilisateur(username, ecole, telephone):
    user = User.objects.create_user(username, password='secret')
    user.profil.role = 'COMPTABLE'
    user.profil.telephone = telephone
    user.profil.ecole = ecole
    user.profil.is_validated = True
    user.profil.save()
    return user


@override_settings(MIDDLEWARE=TEST_MIDDLEWARE)
class HistoriqueAbonnementsTests(TestCase):
    def setUp(self):
        self.today = timezone.localdate()
        self.ecole = Ecole.objects.create(
            nom='École historique', adresse='Conakry',
            telephone='+224620008001', directeur='Direction', etat='VALIDE',
        )
        classe = Classe.objects.create(
            ecole=self.ecole, nom='CE1 A', niveau='PRIMAIRE_2', annee_scolaire='2026-2027',
        )
        self.eleve = Eleve.objects.create(
            matricule='HIS-001', prenom='Moussa', nom='Sylla', sexe='M', classe=classe,
        )
        self.cantine_1 = AbonnementCantine.objects.create(
            eleve=self.eleve, montant=60000, type_repas='COMPLET', periodicite='MENSUEL',
            regime_alimentaire='Sans porc', allergies='Arachides', contact_parent='+224620008010',
            date_debut=self.today - timedelta(days=20), date_expiration=self.today + timedelta(days=10),
            reference_paiement='REC-1',
        )
        self.bus = AbonnementBus.objects.create(
            eleve=self.eleve, montant=150000, periodicite='MENSUEL', zone='Ratoma',
            itineraire='Ligne 2', point_arret='Carrefour Cosa', contact_parent='+224620008011',
            date_debut=self.today - timedelta(days=60), date_expiration=self.today - timedelta(days=30),
        )
        self.user = _utilisateur('historique', self.ecole, '+224620008002')
        self.client.force_login(self.user)

    def test_reprise_cantine_propose_la_suite_du_dernier_abonnement(self):
        data = self.client.get(
            reverse('bus:dernier_abonnement_json', args=[self.eleve.pk, 'cantine'])
        ).json()
        self.assertTrue(data['existe'])
        valeurs = data['valeurs']
        self.assertEqual(valeurs['type_repas'], 'COMPLET')
        self.assertEqual(valeurs['regime_alimentaire'], 'Sans porc')
        self.assertEqual(valeurs['allergies'], 'Arachides')
        self.assertEqual(valeurs['montant'], '60000')
        debut = self.cantine_1.date_expiration + timedelta(days=1)
        self.assertEqual(valeurs['date_debut'], debut.isoformat())
        self.assertGreater(valeurs['date_expiration'], valeurs['date_debut'])
        self.assertNotIn('reference_paiement', valeurs)

    def test_reprise_bus_expire_commence_aujourdhui(self):
        data = self.client.get(
            reverse('bus:dernier_abonnement_json', args=[self.eleve.pk, 'bus'])
        ).json()
        self.assertEqual(data['valeurs']['date_debut'], self.today.isoformat())
        self.assertEqual(data['valeurs']['point_arret'], 'Carrefour Cosa')
        self.assertEqual(data['valeurs']['zone'], 'Ratoma')

    def test_eleve_sans_abonnement(self):
        autre = Eleve.objects.create(
            matricule='HIS-002', prenom='Nene', nom='Bah', sexe='F', classe=self.eleve.classe,
        )
        data = self.client.get(reverse('bus:dernier_abonnement_json', args=[autre.pk, 'bus'])).json()
        self.assertFalse(data['existe'])

    def test_formulaires_preremplis_depuis_lhistorique(self):
        response = self.client.get(reverse('bus:creer_abonnement_cantine'), {'eleve': self.eleve.pk})
        initial = response.context['form'].initial
        self.assertEqual(initial['allergies'], 'Arachides')
        self.assertEqual(initial['date_debut'], self.cantine_1.date_expiration + timedelta(days=1))

        response = self.client.get(reverse('bus:nouveau'), {'eleve': self.eleve.pk})
        self.assertEqual(response.context['form'].initial['point_arret'], 'Carrefour Cosa')

    def test_historique_et_exports(self):
        AbonnementCantine.objects.create(
            eleve=self.eleve, montant=65000, type_repas='COMPLET', periodicite='MENSUEL',
            date_debut=self.cantine_1.date_expiration + timedelta(days=1),
            date_expiration=self.cantine_1.date_expiration + timedelta(days=31),
        )
        page = self.client.get(reverse('bus:historique_eleve', args=[self.eleve.pk, 'cantine']))
        self.assertEqual(page.status_code, 200)
        self.assertEqual(len(page.context['lignes']), 2)
        self.assertEqual(page.context['total'], 125000)
        self.assertContains(page, 'REC-1')

        excel = self.client.get(reverse('bus:historique_eleve_excel', args=[self.eleve.pk, 'cantine']))
        feuille = load_workbook(BytesIO(excel.content)).active
        montants = [feuille.cell(ligne, 4).value for ligne in (5, 6)]
        self.assertEqual(montants, [60000, 65000])

        for nom in ('bus:historique_eleve_pdf', 'bus:carnet_abonnement_pdf'):
            for service in ('bus', 'cantine'):
                with self.subTest(vue=nom, service=service):
                    response = self.client.get(reverse(nom, args=[self.eleve.pk, service]))
                    self.assertEqual(response.status_code, 200)
                    self.assertEqual(response['Content-Type'], 'application/pdf')
                    self.assertTrue(response.content.startswith(b'%PDF'))

    def test_autre_ecole_et_service_inconnu(self):
        self.assertEqual(
            self.client.get(reverse('bus:historique_eleve', args=[self.eleve.pk, 'piscine'])).status_code,
            404,
        )
        autre_ecole = Ecole.objects.create(
            nom='Autre école', adresse='Kindia',
            telephone='+224620008099', directeur='Direction', etat='VALIDE',
        )
        self.client.force_login(_utilisateur('autre-historique', autre_ecole, '+224620008098'))
        for nom in ('bus:historique_eleve', 'bus:carnet_abonnement_pdf', 'bus:dernier_abonnement_json'):
            with self.subTest(vue=nom):
                response = self.client.get(reverse(nom, args=[self.eleve.pk, 'cantine']))
                self.assertEqual(response.status_code, 404)

    def test_libelles_mois_et_jour(self):
        mensuel = AbonnementCantine(date_debut=date(2026, 9, 1), date_expiration=date(2026, 10, 1))
        trimestre = AbonnementCantine(date_debut=date(2026, 9, 1), date_expiration=date(2026, 12, 1))
        self.assertEqual(libelle_mois(mensuel), 'Septembre 2026')
        self.assertEqual(libelle_mois(trimestre), 'Septembre – Décembre 2026')
        self.assertEqual(jour_semaine(date(2026, 9, 25)), 'Vendredi')
